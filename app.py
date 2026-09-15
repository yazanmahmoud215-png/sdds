import hashlib
import os
import re
import secrets
from pathlib import Path

from flask import Flask, flash, redirect, render_template_string, request, session, url_for
from openpyxl import load_workbook


app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
DEFAULT_WORKBOOK = DATA_DIR / "الشهر التاسع.xlsm"

LOGIN_TEMPLATE = """
<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>المحفظة التراكمية</title>
<style>
@import url('https://fonts.googleapis.com/css2?family=Almarai:wght@400;700;800&display=swap');:root{--ink:#17221d;--muted:#637066;--green:#1c6b4d;--cream:#f6f2e9;--gold:#c88d3a;--line:#d8dfd7}*{box-sizing:border-box}body{margin:0;background:var(--cream);color:var(--ink);font-family:'Almarai',Tahoma,sans-serif;min-height:100vh;display:grid;place-items:center;padding:24px}.shell{width:min(940px,100%);display:grid;grid-template-columns:1.05fr .95fr;background:#fff;border:1px solid var(--line);box-shadow:0 24px 70px #1c3b2b18}.intro{padding:clamp(34px,6vw,72px);background:var(--green);color:#fff;position:relative;overflow:hidden}.intro h1{font-size:clamp(30px,4vw,48px);line-height:1.25;margin:42px 0 18px;max-width:390px}.intro p{color:#dce9df;line-height:2;font-size:14px;max-width:350px}.mark{font-size:13px;color:#d7bc83;font-weight:700}.login{padding:clamp(32px,6vw,70px);display:flex;flex-direction:column;justify-content:center}.login h2{margin:0 0 10px;font-size:26px}.hint{color:var(--muted);font-size:13px;line-height:1.8;margin:0 0 30px}label{font-size:13px;font-weight:700;display:block;margin-bottom:9px}input{width:100%;border:1px solid var(--line);padding:15px 16px;font:inherit;font-size:16px;outline:none;direction:ltr;text-align:left;background:#fcfdfb}input:focus{border-color:var(--green);box-shadow:0 0 0 3px #1c6b4d18}button{width:100%;margin-top:18px;border:0;background:var(--gold);color:#fff;padding:15px;font:inherit;font-weight:700;cursor:pointer}.error{background:#fff0ee;color:#9a3d32;border-right:3px solid #c25748;padding:12px;font-size:13px;margin-bottom:18px}@media(max-width:700px){.shell{grid-template-columns:1fr}.intro{padding:30px 28px}.intro h1{margin-top:24px;font-size:32px}.login{padding:32px 28px}}
</style></head><body><main class="shell"><section class="intro"><div class="mark">بيانات خاصة وآمنة</div><h1>المحفظة التراكمية</h1><p>أدخل اسمك للوصول إلى بياناتك الشخصية فقط.</p></section><section class="login"><h2>تسجيل الدخول</h2><p class="hint">اكتب اسمك كما هو موجود في ملف Excel أو اختره من القائمة.</p>{% with messages = get_flashed_messages() %}{% if messages %}<div class="error">{{ messages[0] }}</div>{% endif %}{% endwith %}<form method="post"><label for="name">الاسم</label><input id="name" name="name" list="people" required autocomplete="name" autofocus placeholder="أدخل اسمك هنا"><datalist id="people">{% for name in names %}<option value="{{ name }}">{% endfor %}</datalist><button type="submit">عرض بياناتي</button></form></section></main></body></html>
"""

RESULT_TEMPLATE = """
<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>المحفظة التراكمية</title><style>
@import url('https://fonts.googleapis.com/css2?family=Almarai:wght@400;700;800&display=swap');:root{--ink:#17221d;--green:#1c6b4d;--cream:#f6f2e9;--muted:#637066;--gold:#c88d3a;--line:#d8dfd7}*{box-sizing:border-box}body{margin:0;background:var(--cream);color:var(--ink);font-family:'Almarai',Tahoma,sans-serif;min-height:100vh;padding:30px}.wrap{width:min(900px,100%);margin:auto}.top{display:flex;justify-content:space-between;align-items:center;margin-bottom:42px}.brand{color:var(--green);font-weight:800}.logout{color:var(--muted);font-size:13px;text-decoration:none;border-bottom:1px solid #bbc8bd;padding-bottom:4px}.eyebrow{color:var(--gold);font-size:13px;font-weight:700}.name{font-size:clamp(30px,6vw,56px);margin:13px 0 32px;line-height:1.25}.row-card{background:#fff;border:1px solid var(--line);border-top:5px solid var(--green)}.field{display:grid;grid-template-columns:minmax(210px,38%) 1fr;border-bottom:1px solid var(--line);padding:18px 22px;gap:20px}.field:last-child{border-bottom:0}.label{color:var(--muted);font-size:13px;font-weight:700}.value{font-size:16px;overflow-wrap:anywhere;direction:rtl}.footer{color:var(--muted);font-size:12px;margin-top:24px}@media(max-width:600px){body{padding:22px}.top{margin-bottom:34px}.field{grid-template-columns:1fr;gap:7px;padding:15px 16px}}
</style></head><body><main class="wrap"><div class="top"><div class="brand">المحفظة التراكمية</div><a class="logout" href="{{ url_for('logout') }}">تسجيل الخروج</a></div><div class="eyebrow">الصف الخاص بك من ورقة المحفظة التراكمية</div><h1 class="name">{{ person.name }}</h1><section class="row-card">{% for field in person.fields %}<div class="field"><div class="label">{{ field.label }}</div><div class="value">{{ field.value }}</div></div>{% endfor %}</section><p class="footer">هذه الصفحة تعرض صف اسمك فقط.</p></main></body></html>
"""


def normalize(value):
	text = "" if value is None else str(value).strip().lower()
	text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
	text = text.replace("ى", "ي").replace("ة", "ه")
	text = text.replace("ـ", "")
	text = re.sub(r"[\u064B-\u065F\u0670]", "", text)
	return re.sub(r"[\s_\-\u200c\u200f]+", "", text)


def find_column(headers, aliases):
	aliases = {normalize(alias) for alias in aliases}
	for index, header in enumerate(headers):
		if normalize(header) in aliases:
			return index
	return None


def read_people():
	if not DEFAULT_WORKBOOK.exists():
		return {}
	workbook = load_workbook(DEFAULT_WORKBOOK, read_only=True, data_only=True, keep_vba=False)
	people = {}
	try:
		sheet = workbook["المحفظة التراكمية"] if "المحفظة التراكمية" in workbook.sheetnames else None
		if sheet is None:
			return {}
		rows = sheet.iter_rows(values_only=True)
		headers = next(rows, None)
		if not headers:
			return {}
		name_index = find_column(headers, ["الاسم", "اسم", "name", "الاسم الكامل"])
		if name_index is None:
			return {}
		for row in rows:
			if len(row) <= name_index:
				continue
			name = str(row[name_index]).strip() if row[name_index] is not None else ""
			if not name:
				continue
			fields = []
			for index, header in enumerate(headers):
				if header is None or str(header).strip() == "":
					continue
				value = row[index] if index < len(row) and row[index] is not None else ""
				fields.append({"label": str(header).strip(), "value": value})
			people[normalize(name)] = {"name": name, "fields": fields}
	finally:
		workbook.close()
	return people


def name_fingerprint(name):
	return hashlib.sha256(normalize(name).encode("utf-8")).hexdigest()


@app.route("/", methods=["GET", "POST"])
def login():
	if request.method == "POST":
		name = request.form.get("name", "")
		person = read_people().get(normalize(name))
		if person:
			session.clear()
			session["name"] = name_fingerprint(name)
			return redirect(url_for("profile"))
		flash("الاسم غير موجود.")
	return render_template_string(LOGIN_TEMPLATE, names=[person["name"] for person in read_people().values()])


@app.route("/profile")
def profile():
	name_hash = session.get("name")
	if not name_hash:
		return redirect(url_for("login"))
	person = next((person for name, person in read_people().items() if name_fingerprint(name) == name_hash), None)
	if not person:
		session.clear()
		return redirect(url_for("login"))
	return render_template_string(RESULT_TEMPLATE, person=person)


@app.route("/logout")
def logout():
	session.clear()
	return redirect(url_for("login"))


@app.route("/admin", methods=["GET", "POST"])
def admin():
	admin_key = os.environ.get("ADMIN_KEY")
	if not admin_key or not secrets.compare_digest(request.args.get("key", ""), admin_key):
		return "غير مصرح", 403
	message = ""
	if request.method == "POST":
		uploaded = request.files.get("workbook")
		if not uploaded or not uploaded.filename.lower().endswith((".xlsm", ".xlsx")):
			message = "يرجى اختيار ملف Excel بصيغة xlsm أو xlsx."
		else:
			DEFAULT_WORKBOOK.write_bytes(uploaded.read())
			message = f"تم تحديث ملف البيانات بنجاح. تم تحميل {len(read_people())} شخص."
	return render_template_string("""<!doctype html><html lang='ar' dir='rtl'><meta charset='utf-8'><title>إدارة الملف</title><body style='font-family:Tahoma;max-width:600px;margin:60px auto;padding:20px'><h2>تحديث ملف Excel</h2><p>{{ message }}</p><form method='post' enctype='multipart/form-data'><input type='file' name='workbook' accept='.xlsm,.xlsx' required><button>رفع الملف</button></form></body></html>""", message=message)


if __name__ == "__main__":
	app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)